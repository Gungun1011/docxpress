"""Deterministic test document generator for DocXpress test suite.

Creates test .docx files containing all 12 target element types:
Title, Author, Chapter heading, Subheading, Body paragraphs, Bold/Italic runs,
Table, Figure/image, Caption, Numbered list, Bullet list, References.
Also creates edge-case documents (empty, unicode, interleaved tables).
"""

import io
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


def create_sample_png(file_path: Path) -> Path:
    """Creates a deterministic PNG diagram for embedding as a test figure."""
    img = Image.new("RGB", (320, 160), color=(240, 244, 248))
    draw = ImageDraw.Draw(img)
    
    # Draw simple diagram boxes
    draw.rectangle([20, 30, 130, 130], outline=(41, 128, 185), width=3)
    draw.text((35, 70), "Node A", fill=(41, 128, 185))
    
    draw.line([130, 80, 190, 80], fill=(52, 73, 94), width=3)
    
    draw.rectangle([190, 30, 300, 130], outline=(39, 174, 96), width=3)
    draw.text((205, 70), "Node B", fill=(39, 174, 96))
    
    img.save(file_path, "PNG")
    return file_path


def generate_deterministic_sample(output_path: Path) -> None:
    """Generates the primary deterministic test manuscript containing all 12 element types."""
    doc = docx.Document()
    
    # Metadata
    doc.core_properties.title = "The Architecture of Distributed Systems"
    doc.core_properties.author = "Dr. Ada Lovelace & Alan Turing"
    
    # 1. Title
    p_title = doc.add_paragraph("The Architecture of Distributed Systems", style="Title")
    
    # 2. Author details
    p_author = doc.add_paragraph("Dr. Ada Lovelace & Alan Turing\nDepartment of Computing", style="Subtitle")
    
    # 3. Chapter heading
    p_chap = doc.add_paragraph("Chapter 1: Foundations of Computing", style="Heading 1")
    
    # 4. Subheading
    p_sub = doc.add_paragraph("1.1 The Theoretical Model", style="Heading 2")
    
    # 5. Body paragraph
    doc.add_paragraph(
        "A distributed computational model operates across autonomous nodes communicating via message passing. "
        "Each node maintains local state and coordinates through deterministic consensus algorithms."
    )
    
    # 6. Paragraph with bold, italic, underline, and sized runs
    p_runs = doc.add_paragraph()
    r1 = p_runs.add_run("Typography check: ")
    r2 = p_runs.add_run("bold statement")
    r2.bold = True
    r3 = p_runs.add_run(", followed by an ")
    r4 = p_runs.add_run("italicized caveat")
    r4.italic = True
    r5 = p_runs.add_run(", an ")
    r6 = p_runs.add_run("underlined clause")
    r6.underline = True
    r7 = p_runs.add_run(", and ")
    r8 = p_runs.add_run("large emphasis text")
    r8.font.size = Pt(14)
    r8.font.name = "Georgia"
    r9 = p_runs.add_run(".")
    
    # 7. Table (3x3)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    headers = ["Node ID", "Throughput (ops/s)", "Consensus State"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
        
    data = [
        ["node-01", "12500", "Leader"],
        ["node-02", "11800", "Follower"],
    ]
    for r_idx, row_values in enumerate(data, start=1):
        for c_idx, val in enumerate(row_values):
            table.cell(r_idx, c_idx).text = val
            
    # Spacer
    doc.add_paragraph()
    
    # 8. Figure / Image
    temp_png = output_path.parent / "temp_diagram.png"
    create_sample_png(temp_png)
    doc.add_picture(str(temp_png), width=Inches(3.5))
    if temp_png.exists():
        temp_png.unlink()
        
    # 9. Caption
    p_cap = doc.add_paragraph("Figure 1.1: Distributed consensus state transition between Node A and Node B.")
    p_cap.style = "Caption"
    
    # 10. Numbered list
    doc.add_paragraph("Initialize all consensus registers to state zero.", style="List Number")
    doc.add_paragraph("Broadcast vote request to all active peer nodes.", style="List Number")
    doc.add_paragraph("Commit state transition upon receiving quorum threshold.", style="List Number")
    
    # 11. Bullet list
    doc.add_paragraph("High availability with zero single points of failure.", style="List Bullet")
    doc.add_paragraph("Deterministic log replication with bounded latency.", style="List Bullet")
    doc.add_paragraph("Cryptographic audit trails for state mutations.", style="List Bullet")
    
    # Another Subheading
    doc.add_paragraph("References", style="Heading 2")
    
    # 12. References
    doc.add_paragraph(
        "1. Shannon, C. E. (1948). A Mathematical Theory of Communication. Bell System Technical Journal, 27(3), 379-423."
    )
    doc.add_paragraph(
        "2. Turing, A. M. (1936). On Computable Numbers, with an Application to the Entscheidungsproblem. Proceedings of the London Mathematical Society, 42, 230-265."
    )
    doc.add_paragraph(
        "3. Lamport, L. (1978). Time, Clocks, and the Ordering of Events in a Distributed System. Communications of the ACM, 21(7), 558-565."
    )

    doc.save(str(output_path))


def generate_empty_document(output_path: Path) -> None:
    """Generates an empty DOCX file."""
    doc = docx.Document()
    if doc.paragraphs:
        p = doc.paragraphs[0]._p
        p.getparent().remove(p)
    doc.save(str(output_path))


def generate_unicode_document(output_path: Path) -> None:
    """Generates a DOCX with international Unicode, diacritics, and special characters."""
    doc = docx.Document()
    doc.core_properties.title = "Unicode & Multilingual Test"
    
    doc.add_paragraph("DocXpress: Système d'Édition Décentralisée", style="Title")
    doc.add_paragraph("François René & Müller Schmidt — Zürich, Genève & München")
    doc.add_paragraph("Chapter 1: 普适计算与量子通信", style="Heading 1")
    doc.add_paragraph("В теории информации энтропия определяет меру неопределенности.")
    doc.add_paragraph("Η κβαντική υπολογιστική χρησιμοποιεί qubits για παράλληλο υπολογισμό.")
    doc.add_paragraph("Mathematical symbols: ∀x ∈ S, ∃y : f(x) ≤ y ∧ y → ∞. Emoji check: 🚀 📚 ✨")
    
    doc.save(str(output_path))


def generate_interleaved_document(output_path: Path) -> None:
    """Generates a document with interleaved paragraphs and tables to test document order."""
    doc = docx.Document()
    
    doc.add_paragraph("Paragraph Alpha: Preceding Table 1")
    
    t1 = doc.add_table(rows=2, cols=2)
    t1.cell(0, 0).text = "T1-R1C1"
    t1.cell(0, 1).text = "T1-R1C2"
    t1.cell(1, 0).text = "T1-R2C1"
    t1.cell(1, 1).text = "T1-R2C2"
    
    doc.add_paragraph("Paragraph Beta: Between Table 1 and Table 2")
    
    t2 = doc.add_table(rows=2, cols=1)
    t2.cell(0, 0).text = "T2-R1C1"
    t2.cell(1, 0).text = "T2-R2C1"
    
    doc.add_paragraph("Paragraph Gamma: Following Table 2")
    
    doc.save(str(output_path))


def generate_all_test_documents(output_dir: Path) -> None:
    """Generates all standard test manuscripts into output directory."""
    output_dir.mkdir(parents=True, exist_ok=True)
    
    generate_deterministic_sample(output_dir / "deterministic_sample.docx")
    generate_empty_document(output_dir / "empty.docx")
    generate_unicode_document(output_dir / "unicode_sample.docx")
    generate_interleaved_document(output_dir / "interleaved_tables.docx")


if __name__ == "__main__":
    target_dir = Path(__file__).resolve().parent.parent / "sample_documents"
    print(f"Generating test documents in: {target_dir}")
    generate_all_test_documents(target_dir)
    print("Test documents generated successfully.")
