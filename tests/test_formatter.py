"""Formatting and output-content verification tests."""

from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.parser.docx_parser import DocxParser
from app.services.content_preservation import ContentPreservationError, ContentPreservationService
from app.services.formatter import (
    DEFAULT_PUBLICATION_PROFILE,
    PublicationFormatter,
)


def test_default_formatter_preserves_content_and_applies_profile(deterministic_sample_path: Path, tmp_path: Path):
    output_path = tmp_path / "formatted.docx"
    events = []
    result = PublicationFormatter().format_document(
        deterministic_sample_path,
        output_path,
        progress_callback=events.append,
    )

    assert result.profile_name == DEFAULT_PUBLICATION_PROFILE.name
    assert result.preservation.is_identical is True
    assert result.preservation.source_hash == result.preservation.output_hash
    assert result.preservation.changed_elements == ()
    assert len(events) == result.elements_processed

    output = docx.Document(str(output_path))
    assert abs(output.sections[0].top_margin.cm - 1.52) < 0.01
    assert output.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert output.paragraphs[2].paragraph_format.page_break_before is True
    assert output.paragraphs[4].paragraph_format.line_spacing == 1.5
    assert abs(output.paragraphs[4].paragraph_format.first_line_indent.cm - 1.27) < 0.01


def test_formatter_preserves_tables_images_lists_and_headings(deterministic_sample_path: Path, tmp_path: Path):
    output_path = tmp_path / "formatted.docx"
    PublicationFormatter().format_document(deterministic_sample_path, output_path)
    parsed = DocxParser().parse(output_path)

    source = DocxParser().parse(deterministic_sample_path)
    assert ContentPreservationService.verify_report(source, parsed).is_identical
    assert len(parsed.get_elements_by_type(parsed.elements[6].element_type)) >= 1
    assert len(parsed.get_elements_by_type(parsed.elements[7].element_type)) >= 1
    assert any(element.original_text.startswith("Initialize all") for element in parsed.elements)
    assert any(element.original_text.startswith("Chapter 1") for element in parsed.elements)


def test_formatter_handles_unicode_and_empty_documents(unicode_doc_path: Path, empty_doc_path: Path, tmp_path: Path):
    formatter = PublicationFormatter()
    unicode_output = tmp_path / "unicode.docx"
    empty_output = tmp_path / "empty.docx"

    unicode_result = formatter.format_document(unicode_doc_path, unicode_output)
    empty_result = formatter.format_document(empty_doc_path, empty_output)

    assert unicode_result.preservation.is_identical
    assert empty_result.preservation.is_identical
    assert DocxParser().parse(empty_output).elements == ()


def test_malformed_document_is_rejected(tmp_path: Path):
    malformed = tmp_path / "malformed.docx"
    malformed.write_bytes(b"not a docx")

    with pytest.raises(Exception):
        PublicationFormatter().format_document(malformed, tmp_path / "output.docx")


def test_preservation_report_identifies_changed_missing_and_added_elements(deterministic_sample_path: Path):
    doc = DocxParser().parse(deterministic_sample_path)
    altered = list(doc.elements)
    altered[0] = type(altered[0])(
        element_id=altered[0].element_id,
        element_type=altered[0].element_type,
        paragraph_index=altered[0].paragraph_index,
        original_text="Changed",
        original_style=altered[0].original_style,
        runs=altered[0].runs,
    )
    changed = doc.with_elements(altered[:-1])
    report = ContentPreservationService.verify_report(doc, changed)

    assert report.is_identical is False
    assert report.changed_elements == (0,)
    assert report.missing_elements == (len(doc.elements) - 1,)
    assert report.added_elements == ()